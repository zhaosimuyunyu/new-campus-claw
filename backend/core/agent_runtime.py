import asyncio
from contextvars import ContextVar
import base64
import hashlib
import html
import json
import mimetypes
import os
import re
import inspect
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import httpx
import fitz
from jsonschema import Draft202012Validator
from docx import Document
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from django.conf import settings
from django.core.files.base import ContentFile
from django.db import connection
from django.utils import timezone
from pgvector.django import CosineDistance
from asgiref.sync import sync_to_async

from .models import (AgentArtifact, AgentAttachment, AgentEvent, AgentMessage, AgentRun, AgentThread,
                     ChunkEmbedding, KnowledgeChunk, KnowledgeEntry, Material, Membership)


class RuntimeErrorWithStatus(Exception):
    def __init__(self, message, status=503):
        self.status = status
        super().__init__(message)


def split_chunks(text, strategy="auto", max_chars=800, overlap=80, separators=None):
    if strategy not in {"auto", "custom", "hierarchy"}:
        raise ValueError("不支持的切片策略")
    source = text
    if strategy == "hierarchy":
        sections = re.split(r"(?=^#{1,3}\s+)", source, flags=re.M)
        spans = []
        cursor = 0
        for section in sections:
            if not section:
                continue
            pos = source.find(section, cursor)
            spans.append((section, pos))
            cursor = pos + len(section)
        windows = []
        for section, start in spans:
            if len(section) <= 800:
                windows.append((section.strip(), start + len(section) - len(section.lstrip())))
            else:
                windows.extend(_window_chunks(section, start, 800, 80, ["\n\n", "\n", "。", ". "]))
    else:
        if strategy == "auto":
            max_chars, overlap, separators = 800, 80, ["\n\n", "\n", "。", ". "]
        else:
            max_chars = max(100, min(int(max_chars), 2000))
            overlap = int(max_chars * max(0, min(float(overlap), 0.5))) if float(overlap) <= 0.5 else int(overlap)
            overlap = max(0, min(overlap, max_chars // 2))
            separators = separators or ["\n\n", "\n", "。", ". "]
        windows = _window_chunks(source, 0, max_chars, overlap, separators)
    return [(part, start, start + len(part)) for part, start in windows if part.strip()]


def _window_chunks(text, base, size, overlap, separators):
    chunks = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        if end < n:
            floor = start + size // 2
            breaks = [text.rfind(sep, floor, end) for sep in separators if sep]
            point = max(breaks, default=-1)
            if point > start:
                end = point + 1
        left = start
        right = end
        while left < right and text[left].isspace():
            left += 1
        while right > left and text[right - 1].isspace():
            right -= 1
        if left < right:
            chunks.append((text[left:right], base + left))
        if end >= n:
            break
        start = max(start + 1, end - overlap)
    return chunks


def extract_text(file_path, mime_type=""):
    suffix = Path(file_path).suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        return Path(file_path).read_text(encoding="utf-8", errors="replace"), suffix[1:]
    if suffix == ".pdf":
        reader = PdfReader(file_path)
        return _valid_unicode("\n\n".join(page.extract_text() or "" for page in reader.pages)), "pdf"
    if suffix == ".docx":
        doc = Document(file_path)
        return _valid_unicode("\n".join(p.text for p in doc.paragraphs)), "docx"
    if mime_type.startswith("image/"):
        return "", "image"
    raise ValueError("此文件格式暂不支持文本索引")


def _valid_unicode(value):
    return value.encode("utf-8", errors="replace").decode("utf-8")


def _openai_headers(key):
    return {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}


async def embed_texts(texts):
    base = (os.getenv("EMBEDDING_BASE_URL") or os.getenv("EMBEDDING_API_BASE", "")).rstrip("/")
    key = os.getenv("EMBEDDING_API_KEY") or os.getenv("MODEL_API_KEY", "")
    model = os.getenv("EMBEDDING_MODEL", "")
    if not base or not model or not key:
        raise RuntimeErrorWithStatus("未配置嵌入模型服务")
    async with httpx.AsyncClient(timeout=60) as client:
        response = await client.post(base + "/embeddings", headers=_openai_headers(key), json={"model": model, "input": texts})
        response.raise_for_status()
    vectors = [row["embedding"] for row in response.json().get("data", [])]
    expected = int(os.getenv("EMBEDDING_DIM") or os.getenv("EMBEDDING_DIMENSIONS") or "2048")
    if len(vectors) != len(texts) or any(len(v) != expected for v in vectors):
        raise RuntimeErrorWithStatus(f"嵌入模型返回维度不符，要求 {expected} 维")
    return vectors, model


def index_material(material, strategy="auto", max_chars=800, overlap=80, separators=None):
    entry, _ = KnowledgeEntry.objects.get_or_create(material=material)
    entry.status = KnowledgeEntry.Status.PENDING
    entry.error_message = ""
    entry.save(update_fields=["status", "error_message"])
    try:
        with material.attachment.file.open("rb") as source:
            raw = source.read()
        suffix = Path(material.attachment.original_name).suffix.lower()
        temp = settings.MEDIA_ROOT / f".extract-{material.id}{suffix}"
        temp.parent.mkdir(parents=True, exist_ok=True)
        temp.write_bytes(raw)
        try:
            body, parser = extract_text(temp, mimetypes.guess_type(material.attachment.original_name)[0] or "")
        finally:
            temp.unlink(missing_ok=True)
        if not body.strip():
            raise ValueError("文件中没有可索引的文本")
        pieces = split_chunks(body, strategy, max_chars, overlap, separators)
        if not pieces:
            raise ValueError("未能从资料生成切片")
        from asgiref.sync import async_to_sync
        vectors, model = async_to_sync(embed_texts)([piece[0] for piece in pieces])
        entry.chunks.all().delete()
        entry.body_text = body
        entry.parser = parser
        entry.content_hash = hashlib.sha256(raw).hexdigest()
        entry.chunk_strategy = strategy
        entry.status = KnowledgeEntry.Status.READY
        entry.indexed_at = timezone.now()
        entry.save()
        for index, ((chunk_text, start, end), vector) in enumerate(zip(pieces, vectors)):
            chunk = KnowledgeChunk.objects.create(entry=entry, chunk_index=index, chunk_text=chunk_text,
                start_offset=start, end_offset=end, status=KnowledgeChunk.Status.READY)
            ChunkEmbedding.objects.create(chunk=chunk, embedding=vector, model=model, dimensions=len(vector))
    except Exception as exc:
        entry.status = KnowledgeEntry.Status.FAILED
        entry.error_message = str(exc)[:2000]
        entry.save(update_fields=["status", "error_message"])
        raise
    return entry


def require_course_member(user, course_id):
    membership = Membership.objects.filter(course_id=course_id, user=user).select_related("course").first()
    return membership


def keyword_search(course_id, query, limit=10):
    if not query.strip():
        return []
    with connection.cursor() as cursor:
        cursor.execute("SELECT chunk.id FROM core_knowledgechunk chunk JOIN core_knowledgeentry entry ON entry.id=chunk.entry_id JOIN core_material material ON material.id=entry.material_id WHERE material.course_id=%s AND entry.status='ready' AND chunk.status='ready' AND chunk.chunk_text %% %s ORDER BY similarity(chunk.chunk_text,%s) DESC LIMIT %s", [course_id, query, query, limit])
        ids = [row[0] for row in cursor.fetchall()]
    chunks = {chunk.id: chunk for chunk in KnowledgeChunk.objects.filter(id__in=ids).select_related("entry__material")}
    return [serialize_chunk(chunks[i], "keyword") for i in ids if i in chunks]


async def vector_search(course_id, query, limit=10):
    vectors, _ = await embed_texts([query])
    return await asyncio.to_thread(_vector_search_sync, course_id, vectors[0], limit)


def _vector_search_sync(course_id, vector, limit):
    rows = list(ChunkEmbedding.objects.filter(chunk__entry__material__course_id=course_id,
        chunk__entry__status="ready", chunk__status="ready", embedding__isnull=False)
        .annotate(distance=CosineDistance("embedding", vector)).order_by("distance")
        .select_related("chunk__entry__material")[:limit * 3])
    hits = []
    for row in rows:
        score = 1 - row.distance
        if score < 0.35:
            continue
        hit = serialize_chunk(row.chunk, "vector")
        hit["score"] = score
        hits.append(hit)
        if len(hits) >= limit:
            break
    return hits


def serialize_chunk(chunk, mode):
    material = chunk.entry.material
    return {"chunk_id": chunk.id, "material_id": material.id, "attachment_id": material.attachment_id, "title": material.title,
            "chunk_index": chunk.chunk_index, "start_offset": chunk.start_offset, "end_offset": chunk.end_offset,
            "excerpt": chunk.chunk_text, "mode": mode, "file_name": material.attachment.original_name,
            "download_url": f"/api/files/{material.attachment_id}/download/",
            "preview_url": f"/api/materials/{material.id}/preview/"}


async def search(course_id, query, mode="hybrid", limit=10):
    if not query.strip():
        raise ValueError("检索内容不能为空")
    if mode not in {"keyword", "vector", "hybrid"}:
        raise ValueError("不支持的检索模式")
    if mode == "keyword":
        return await asyncio.to_thread(keyword_search, course_id, query, limit)
    if mode == "vector":
        return await vector_search(course_id, query, limit)
    keyword_hits, vector_hits = await asyncio.gather(
        asyncio.to_thread(keyword_search, course_id, query, limit * 3), vector_search(course_id, query, limit * 3))
    fused = {}
    for hits in (keyword_hits, vector_hits):
        for rank, hit in enumerate(hits, 1):
            item = fused.setdefault(hit["chunk_id"], {**hit, "rrf_score": 0.0})
            item["rrf_score"] += 1 / (60 + rank)
            item["mode"] = "hybrid"
    return sorted(fused.values(), key=lambda item: item["rrf_score"], reverse=True)[:limit]


def event(run, event_type, payload):
    sequence = (run.events.order_by("-sequence").values_list("sequence", flat=True).first() or 0) + 1
    item = AgentEvent.objects.create(run=run, event_type=event_type, payload=payload, sequence=sequence)
    return item


@dataclass
class ToolSpec:
    name: str
    version: str
    input_schema: dict
    output_schema: dict
    required_permission: str
    accesses_course_data: bool
    has_file_side_effect: bool


@dataclass
class ToolContext:
    user: object
    thread: AgentThread | None = None
    permissions: frozenset[str] = frozenset()


class ModelAdapter:
    """Stable contract for streaming chat model gateways."""
    async def stream(self, messages, attachments=()):
        raise NotImplementedError


class EmbeddingAdapter:
    """Stable contract for embedding providers used by the indexer."""
    async def embed(self, texts):
        raise NotImplementedError


class VisionAdapter:
    """Stable contract for image and scanned-page understanding."""
    async def understand(self, parts, prompt=""):
        raise NotImplementedError


@dataclass(frozen=True)
class SkillManifest:
    name: str
    version: str
    prompt_template: str = ""
    tool_dependencies: tuple[str, ...] = ()
    input_schema: dict | None = None
    output_schema: dict | None = None


class MCPToolProvider:
    async def list_tools(self):
        raise NotImplementedError

    async def call_tool(self, name, arguments, context: ToolContext):
        raise NotImplementedError


class HarnessRunner:
    async def run(self, scenario, context):
        raise NotImplementedError


class ArtifactStore:
    def save(self, owner, thread, title, body, kind):
        return create_artifact(owner, thread, title, body, kind)


class ToolRegistry:
    def __init__(self):
        self._tools = {}

    def register(self, spec, function):
        if spec.name in self._tools:
            raise ValueError(f"工具已注册: {spec.name}")
        self._tools[spec.name] = (spec, function)

    def describe(self):
        return [spec for spec, _ in self._tools.values()]

    def get(self, name):
        return self._tools[name]

    async def execute(self, name, arguments, context: ToolContext):
        spec, function = self.get(name)
        if not isinstance(arguments, dict):
            raise ValueError("工具参数必须是 JSON 对象")
        if spec.required_permission not in context.permissions:
            raise PermissionError("当前运行没有该工具权限")
        errors = list(Draft202012Validator(spec.input_schema).iter_errors(arguments))
        if errors:
            raise ValueError("工具参数无效: " + errors[0].message)
        if function is None:
            raise RuntimeError("工具尚未绑定执行器")
        result = function(arguments, context)
        result = await result if inspect.isawaitable(result) else result
        if not Draft202012Validator(spec.output_schema).is_valid(result):
            raise ValueError("工具返回值不符合声明格式")
        return result


def builtin_tool_registry():
    registry = ToolRegistry()

    async def search_course(arguments, context):
        if not context.thread or not context.thread.course_id:
            raise ValueError("工具需要课程线程")
        mode = arguments.get("mode", "hybrid")
        return await search(context.thread.course_id, str(arguments.get("query", "")), mode, 10)

    async def list_materials(arguments, context):
        if not context.thread or not context.thread.course_id:
            raise ValueError("工具需要课程线程")
        rows = await sync_to_async(list)(Material.objects.filter(course_id=context.thread.course_id).select_related("attachment").order_by("title"))
        return [serialize_material(item) for item in rows]

    async def make_artifact(kind, arguments, context):
        if not context.thread:
            raise ValueError("工具需要线程")
        item = await sync_to_async(create_artifact)(context.user, context.thread, arguments.get("title"), str(arguments.get("content", "")), kind)
        return {"id": item.id, "name": item.name, "mime_type": item.mime_type,
                "download_url": f"/api/agent/artifacts/{item.id}/download/"}

    def artifact_tool(kind):
        async def execute(arguments, context):
            return await make_artifact(kind, arguments, context)
        return execute

    registry.register(ToolSpec("search_course_knowledge", "1.0",
        {"type": "object", "properties": {"query": {"type": "string"}, "mode": {"type": "string", "enum": ["keyword", "vector", "hybrid"]}}, "required": ["query"], "additionalProperties": False},
        {"type": "array"}, "course.read", True, False), search_course)
    registry.register(ToolSpec("list_course_materials", "1.0", {"type": "object", "properties": {}, "additionalProperties": False}, {"type": "array"}, "course.read", True, False), list_materials)
    artifact_schema = {"type": "object", "properties": {"title": {"type": "string"}, "content": {"type": "string"}}, "required": ["content"], "additionalProperties": False}
    registry.register(ToolSpec("create_pdf", "1.0", artifact_schema, {"type": "object"}, "artifact.create", False, True), artifact_tool("pdf"))
    registry.register(ToolSpec("create_docx", "1.0", artifact_schema, {"type": "object"}, "artifact.create", False, True), artifact_tool("docx"))
    return registry


def model_tools(registry, permissions):
    return [{"type": "function", "function": {"name": spec.name,
             "description": f"CampusClaw 内置工具，版本 {spec.version}", "parameters": spec.input_schema}}
            for spec in registry.describe() if spec.required_permission in permissions]


def serialize_material(material):
    return {"id": material.id, "title": material.title, "name": material.attachment.original_name,
            "size": material.attachment.size, "download_url": f"/api/files/{material.attachment_id}/download/"}


def artifact_html(title, body):
    escaped = html.escape(body).replace("\n", "<br>")
    return f"<!doctype html><html lang='zh-CN'><meta charset='utf-8'><title>{html.escape(title)}</title><body><h1>{html.escape(title)}</h1><p>{escaped}</p></body></html>"


def create_artifact(owner, thread, title, body, kind):
    title = (title or "对话生成文档").strip()[:160]
    if kind == "docx":
        doc = Document()
        doc.add_heading(title, 0)
        for paragraph in body.split("\n"):
            doc.add_paragraph(paragraph)
        stream = BytesIO()
        doc.save(stream)
        content = stream.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = f"{title}.docx"
    elif kind == "pdf":
        document = fitz.open()
        page = document.new_page(width=595, height=842)
        fontfile = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
        font_kwargs = {"fontname": "CampusClawCJK"} if os.path.exists(fontfile) else {}
        if font_kwargs:
            page.insert_font(fontname="CampusClawCJK", fontfile=fontfile)
        page.insert_text((50, 65), title, fontsize=20, color=(0.09, 0.25, 0.22), **font_kwargs)
        box = fitz.Rect(50, 95, 545, 790)
        page.insert_textbox(box, body, fontsize=11, lineheight=1.5, color=(0.12, 0.16, 0.16), **font_kwargs)
        content = document.tobytes()
        document.close()
        mime = "application/pdf"
        filename = f"{title}.pdf"
    else:
        raise ValueError("仅支持 PDF 与 DOCX")
    artifact = AgentArtifact(owner=owner, thread=thread, course_id=thread.course_id, name=filename, mime_type=mime,
                             preview_html=artifact_html(title, body))
    artifact.file.save(filename, ContentFile(content), save=True)
    return artifact


trace_sink = ContextVar("model_trace_sink", default=None)
trace_frames = ContextVar("model_trace_frames", default=None)


async def trace_record(kind, payload):
    sink = trace_sink.get()
    if sink:
        await sink(kind, payload)


async def model_stream(messages, attachments=(), tools=()):
    frames = []
    token = trace_frames.set(frames)
    text_parts, calls = [], []
    try:
        async for delta, model in _model_stream_impl(messages, attachments, tools):
            if isinstance(delta, dict):
                calls.extend(delta.get("tool_calls", []))
            else:
                text_parts.append(delta)
            yield delta, model
    finally:
        await trace_record("llm_response", {"content": "".join(text_parts), "tool_calls": calls, "frames": frames})
        trace_frames.reset(token)


async def _model_stream_impl(messages, attachments=(), tools=()):
    provider = os.getenv("MODEL_PROVIDER", "mock").lower()
    model = os.getenv("MODEL_NAME", "")
    if provider == "mock" or not model:
        await trace_record("llm_request", {"provider": "mock", "endpoint": None, "sent": False,
            "payload": {"model": model or "mock-local", "messages": messages, "tools": tools, "stream": True}})
        answer = "【演示模式，未调用真实模型】我已收到你的问题。配置课程模型网关后，我可以继续结合课程资料和附件生成回答。"
        for token in re.findall(r".{1,8}", answer, flags=re.S):
            await asyncio.sleep(0.02)
            yield token, model or "mock-local"
        return
    if provider == "anthropic":
        base = os.getenv("MODEL_API_BASE", "https://api.anthropic.com").rstrip("/")
        key = os.getenv("ANTHROPIC_API_KEY") or os.getenv("MODEL_API_KEY", "")
        if not key:
            raise RuntimeErrorWithStatus("未配置 Anthropic API 密钥")
        endpoint = base + "/v1/messages"
        anthropic_messages = []
        for row in messages:
            if row["role"] == "system":
                continue
            content = row["content"]
            if isinstance(content, list):
                blocks = []
                for part in content:
                    if part.get("type") == "image_url":
                        header, encoded = part["image_url"]["url"].split(",", 1)
                        media_type = header.split(":", 1)[1].split(";", 1)[0]
                        blocks.append({"type": "image", "source": {"type": "base64", "media_type": media_type, "data": encoded}})
                    else:
                        blocks.append(part)
                content = blocks
            if row["role"] == "tool":
                anthropic_messages.append({"role": "user", "content": [{"type": "tool_result", "tool_use_id": row["tool_call_id"], "content": content}]})
                continue
            if row.get("tool_calls"):
                content = ([{"type": "text", "text": content}] if content else []) + [
                    {"type": "tool_use", "id": call["id"], "name": call["function"]["name"], "input": json.loads(call["function"]["arguments"] or "{}")} for call in row["tool_calls"]]
            anthropic_messages.append({"role": row["role"], "content": content})
        payload = {"model": model, "max_tokens": 2048, "stream": True,
                   "messages": anthropic_messages,
                   "system": next((row["content"] for row in messages if row["role"] == "system"), "")}
        if tools:
            payload["tools"] = [{"name": t["function"]["name"], "description": t["function"]["description"], "input_schema": t["function"]["parameters"]} for t in tools]
        headers = {"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"}
        parser = "anthropic"
    else:
        base = os.getenv("MODEL_API_BASE", "").rstrip("/")
        key = os.getenv("MODEL_API_KEY", "")
        if not base or not key:
            raise RuntimeErrorWithStatus("未配置 OpenAI 兼容模型网关")
        endpoint = base + "/chat/completions"
        payload = {"model": model, "stream": True, "messages": messages}
        if tools:
            payload["tools"] = tools
        headers = _openai_headers(key)
        parser = "openai"
    await trace_record("llm_request", {"provider": provider, "endpoint": endpoint, "sent": True, "payload": payload})
    async with httpx.AsyncClient(timeout=httpx.Timeout(180, connect=20)) as client:
        async with client.stream("POST", endpoint, headers=headers, json=payload) as response:
            if response.status_code >= 400:
                detail = (await response.aread()).decode("utf-8", "replace")
                raise RuntimeErrorWithStatus(f"模型网关请求失败: HTTP {response.status_code} {detail[:300]}")
            tool_calls = {}
            stream_finished = False
            async for line in response.aiter_lines():
                if not line.startswith("data:"):
                    continue
                raw = line[5:].strip()
                if raw == "[DONE]":
                    stream_finished = True
                    break
                try:
                    data = json.loads(raw)
                except json.JSONDecodeError:
                    continue
                if trace_frames.get() is not None:
                    trace_frames.get().append(data)
                if parser == "anthropic":
                    if data.get("type") == "error":
                        raise RuntimeErrorWithStatus("模型流式生成失败")
                    if data.get("delta", {}).get("stop_reason") == "max_tokens":
                        raise RuntimeErrorWithStatus("回答达到长度上限，尚未完成，请要求模型缩短页面后重试")
                    if data.get("type") == "message_stop":
                        stream_finished = True
                    if data.get("type") == "content_block_start" and data.get("content_block", {}).get("type") == "tool_use":
                        block = data["content_block"]
                        tool_calls[data["index"]] = {"id": block["id"], "type": "function", "function": {"name": block["name"], "arguments": ""}}
                    if data.get("type") == "content_block_delta" and data.get("delta", {}).get("type") == "input_json_delta":
                        tool_calls[data["index"]]["function"]["arguments"] += data["delta"].get("partial_json", "")
                    if data.get("type") == "content_block_delta":
                        yield data.get("delta", {}).get("text", ""), model
                else:
                    if data.get("error"):
                        raise RuntimeErrorWithStatus("模型流式生成失败")
                    choice = (data.get("choices") or [{}])[0]
                    finish_reason = choice.get("finish_reason")
                    if finish_reason in {"length", "content_filter"}:
                        raise RuntimeErrorWithStatus("回答被截断，尚未完成，请缩短请求后重试")
                    if finish_reason:
                        stream_finished = True
                    delta_data = choice.get("delta", {})
                    for call in delta_data.get("tool_calls") or []:
                        item = tool_calls.setdefault(call["index"], {"id": "", "type": "function", "function": {"name": "", "arguments": ""}})
                        item["id"] += call.get("id") or ""
                        function = call.get("function") or {}
                        item["function"]["name"] += function.get("name") or ""
                        item["function"]["arguments"] += function.get("arguments") or ""
                    delta = delta_data.get("content", "")
                    if isinstance(delta, str) and delta:
                        yield delta, model
            if not stream_finished:
                raise RuntimeErrorWithStatus("模型连接提前结束，回答尚未完成，请重试")
            if tool_calls:
                yield {"tool_calls": [tool_calls[index] for index in sorted(tool_calls)]}, model


def attachment_prompt(attachment):
    mime = attachment.mime_type
    if mime.startswith("image/"):
        data = base64.b64encode(attachment.file.read()).decode("ascii")
        return {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{data}"}}
    if Path(attachment.original_name).suffix.lower() == ".pdf":
        with fitz.open(attachment.file.path) as document:
            pages = []
            extracted = []
            for page_number, page in enumerate(document):
                text = _valid_unicode(page.get_text()).strip()
                if text:
                    extracted.append(text)
                elif page_number < 5:
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                    encoded = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
                    pages.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{encoded}"}})
                if page_number >= 4:
                    break
            content = [{"type": "text", "text": f"PDF 附件 {attachment.original_name}，最多处理前 5 页。\n" + "\n".join(extracted)[:30000]}]
            content.extend(pages)
            return {"type": "text", "text": content[0]["text"]} if len(content) == 1 else content
    text, _ = extract_text(attachment.file.path, mime)
    return {"type": "text", "text": f"附件 {attachment.original_name}:\n{text[:30000]}"}
